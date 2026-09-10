"""`scripts/backup_tuzilma.py` — zaxira "tushunarli tuzilma" (Rasm/Excel/Nakladnoy).

Bazadan real o'qiydi (db fixture), STORAGE_PATH monkeypatch bilan tmp_path'ga
yo'naltiriladi. Real Telegram/kamera bu yerda ishtirok etmaydi.
"""

import uuid
from datetime import date, datetime, timezone
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select

from app.core.config import settings
from app.models.kip import Kip, KipHolati
from app.models.mahsulot import Mahsulot
from app.models.partiya import Partiya, PartiyaHolati
from scripts import backup_tuzilma


@pytest.fixture()
def mahsulot_lint(db) -> Mahsulot:
    m = Mahsulot(kod="lint", nomi="Lint")
    db.add(m)
    db.commit()
    db.refresh(m)
    return m


def _partiya(db, mahsulot_id: int, raqami: int) -> Partiya:
    p = Partiya(mahsulot_id=mahsulot_id, partiya_raqami=raqami, holati=PartiyaHolati.ochiq)
    db.add(p)
    db.commit()
    db.refresh(p)
    return p


def _kip(
    db,
    *,
    partiya: Partiya,
    operator_id: int,
    kip_raqami: int,
    ogirlik: float,
    smena,
    vaqt: datetime,
    surat_yoli: str | None = None,
    holati: KipHolati = KipHolati.aktiv,
) -> Kip:
    k = Kip(
        mijoz_id=str(uuid.uuid4()),
        partiya_id=partiya.id,
        kip_raqami=kip_raqami,
        ogirlik=ogirlik,
        smena=smena,
        operator_id=operator_id,
        mahalliy_vaqt=vaqt,
        vaqt=vaqt,
        surat_yoli=surat_yoli,
        holati=holati,
    )
    db.add(k)
    db.commit()
    db.refresh(k)
    return k


def _kun(db, kip_id: int) -> date:
    """Bazaning `func.date(vaqt)` natijasi (skript ham aynan shundan foydalanadi)."""
    return backup_tuzilma._kun(
        db.execute(select(func.date(Kip.vaqt)).where(Kip.id == kip_id)).scalar()
    )


def _surat_yoz(storage: Path, nisbiy: str, mazmun: bytes = b"jpeg-bytes") -> None:
    p = storage / nisbiy
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_bytes(mazmun)


# --------------------------------------------------------------------------


def test_excel_kombinatsiyalari_faqat_real_kunlik_smena_mahsulot(
    db, operator, mahsulot_tola, mahsulot_lint
):
    from app.models.foydalanuvchi import Smena

    p_tola = _partiya(db, mahsulot_tola.id, 10)
    p_lint = _partiya(db, mahsulot_lint.id, 20)
    v1 = datetime(2026, 9, 8, 12, 0, tzinfo=timezone.utc)
    v2 = datetime(2026, 9, 9, 12, 0, tzinfo=timezone.utc)

    _kip(db, partiya=p_tola, operator_id=operator.id, kip_raqami=1, ogirlik=100, smena=Smena.A, vaqt=v1)
    _kip(db, partiya=p_tola, operator_id=operator.id, kip_raqami=2, ogirlik=110, smena=Smena.A, vaqt=v1)
    _kip(db, partiya=p_lint, operator_id=operator.id, kip_raqami=1, ogirlik=90, smena=Smena.B, vaqt=v1)
    _kip(db, partiya=p_tola, operator_id=operator.id, kip_raqami=3, ogirlik=120, smena=Smena.A, vaqt=v2)
    # bekor qilingan -> kombinatsiya BERMASLIGI kerak
    _kip(
        db, partiya=p_lint, operator_id=operator.id, kip_raqami=2, ogirlik=80,
        smena=Smena.C, vaqt=v2, holati=KipHolati.bekor_qilingan,
    )

    kombinatsiyalar = backup_tuzilma.excel_kombinatsiyalari(db)
    sodda = {(k.isoformat(), s.value, mid) for k, s, mid in kombinatsiyalar}

    k1 = _kip(db, partiya=p_tola, operator_id=operator.id, kip_raqami=9, ogirlik=1, smena=Smena.A, vaqt=v1).id
    kun1 = _kun(db, k1).isoformat()

    assert (kun1, "A", mahsulot_tola.id) in sodda
    assert (kun1, "B", mahsulot_lint.id) in sodda
    # bekor qilingan yagona C-smena kipi edi -> C umuman yo'q
    assert not any(s == "C" for _, s, _ in sodda)


def test_suratlarni_joylashtir_tuzilma_va_nom(
    db, operator, mahsulot_tola, mahsulot_lint, tmp_path, monkeypatch
):
    from app.models.foydalanuvchi import Smena

    storage = tmp_path / "storage"
    dest = tmp_path / "storage_2026-09-09_2000"
    monkeypatch.setattr(settings, "STORAGE_PATH", str(storage))

    p_tola = _partiya(db, mahsulot_tola.id, 55)
    v = datetime(2026, 9, 9, 12, 0, tzinfo=timezone.utc)

    _surat_yoz(storage, "2026-09/2026-09-09/Smena_A/Tola/aaaa.jpg")
    _surat_yoz(storage, "2026-09/2026-09-09/Smena_A/Tola/bbbb.jpg")
    k1 = _kip(db, partiya=p_tola, operator_id=operator.id, kip_raqami=4, ogirlik=142.6,
              smena=Smena.A, vaqt=v, surat_yoli="2026-09/2026-09-09/Smena_A/Tola/aaaa.jpg")
    # surati bazada ko'rsatilgan, lekin diskda YO'Q -> "topilmadi", crash emas
    _kip(db, partiya=p_tola, operator_id=operator.id, kip_raqami=5, ogirlik=100,
         smena=Smena.A, vaqt=v, surat_yoli="2026-09/2026-09-09/Smena_A/Tola/yoq.jpg")
    # bekor qilingan -> "KIP-Tarozi Rasm"ga TUSHMASLIGI kerak
    _kip(db, partiya=p_tola, operator_id=operator.id, kip_raqami=6, ogirlik=99,
         smena=Smena.A, vaqt=v, surat_yoli="2026-09/2026-09-09/Smena_A/Tola/bbbb.jpg",
         holati=KipHolati.bekor_qilingan)

    natija = backup_tuzilma.suratlarni_joylashtir(db, dest, storage)

    kun = _kun(db, k1.id).strftime("%d.%m.%Y")
    kutilgan = dest / "KIP-Tarozi Rasm" / "Sentabr" / kun / "Smena_A" / "Tola" / "Tola_Partiya55_Kip4_142.6kg.jpg"
    assert kutilgan.is_file()
    assert natija["nusxalandi"] == 1
    assert natija["topilmadi"] == 1
    # bekor qilingan kip surati nusxalanmadi
    assert not (dest / "KIP-Tarozi Rasm" / "Sentabr" / kun / "Smena_A" / "Tola" / "Tola_Partiya55_Kip6_99kg.jpg").exists()


def test_excellarni_yasa_fayl_va_mazmun(
    db, operator, mahsulot_tola, tmp_path, monkeypatch
):
    from app.models.foydalanuvchi import Smena

    storage = tmp_path / "storage"
    dest = tmp_path / "storage_2026-09-09_2000"
    monkeypatch.setattr(settings, "STORAGE_PATH", str(storage))

    p_tola = _partiya(db, mahsulot_tola.id, 7)
    v = datetime(2026, 9, 9, 12, 0, tzinfo=timezone.utc)
    k1 = _kip(db, partiya=p_tola, operator_id=operator.id, kip_raqami=1, ogirlik=100, smena=Smena.A, vaqt=v)
    _kip(db, partiya=p_tola, operator_id=operator.id, kip_raqami=2, ogirlik=150.5, smena=Smena.A, vaqt=v)

    natija = backup_tuzilma.excellarni_yasa(db, dest)
    assert natija["yasaldi"] == 1
    assert natija["xato"] == 0

    kun = _kun(db, k1.id)
    papka = dest / "KIP-Tarozi Excel" / "Sentabr" / kun.strftime("%d.%m.%Y") / "Smena_A" / "Tola"
    fayl = papka / f"Smena_A_Tola_{kun.isoformat()}.xlsx"
    assert fayl.is_file()

    ws = load_workbook(fayl).active
    jami_qatori = [r for r in ws.iter_rows(values_only=True) if r[0] and str(r[0]).startswith("JAMI")][0]
    assert jami_qatori[3] == pytest.approx(250.5)


def test_nakladnoylarni_kochir(tmp_path):
    storage = tmp_path / "storage"
    dest = tmp_path / "storage_x"
    (storage / "nakladnoy").mkdir(parents=True)
    (storage / "nakladnoy" / "N-001.pdf").write_bytes(b"%PDF-1.4")
    (storage / "nakladnoy" / "N-002.pdf").write_bytes(b"%PDF-1.4")
    (storage / "nakladnoy" / "boshqa.txt").write_text("x")

    natija = backup_tuzilma.nakladnoylarni_kochir(dest, storage)
    assert natija["nusxalandi"] == 2
    assert (dest / "KIP-Tarozi Nakladnoy" / "N-001.pdf").is_file()
    assert not (dest / "KIP-Tarozi Nakladnoy" / "boshqa.txt").exists()


def test_nakladnoy_papka_yoq_bolsa_xato_bermaydi(tmp_path):
    natija = backup_tuzilma.nakladnoylarni_kochir(tmp_path / "d", tmp_path / "bosh_storage")
    assert natija == {"nusxalandi": 0, "xato": 0}


def test_tuzilma_yasa_smoke(db, operator, mahsulot_tola, tmp_path, monkeypatch):
    """`tuzilma_yasa()` `app.core.database.SessionLocal` orqali o'z sessiyasini
    ochadi — testda uni shu testning tranzaksiyasiga yo'naltiramiz."""
    from sqlalchemy.orm import sessionmaker

    from app.models.foydalanuvchi import Smena

    storage = tmp_path / "storage"
    dest = tmp_path / "storage_2026-09-09_2000"
    monkeypatch.setattr(settings, "STORAGE_PATH", str(storage))
    monkeypatch.setattr("app.core.database.SessionLocal", sessionmaker(bind=db.connection()))

    p = _partiya(db, mahsulot_tola.id, 1)
    v = datetime(2026, 9, 9, 12, 0, tzinfo=timezone.utc)
    _surat_yoz(storage, "2026-09/2026-09-09/Smena_A/Tola/z.jpg")
    _kip(db, partiya=p, operator_id=operator.id, kip_raqami=1, ogirlik=120, smena=Smena.A, vaqt=v,
         surat_yoli="2026-09/2026-09-09/Smena_A/Tola/z.jpg")
    (storage / "nakladnoy").mkdir(parents=True)
    (storage / "nakladnoy" / "N-9.pdf").write_bytes(b"%PDF")

    xulosa = backup_tuzilma.tuzilma_yasa(dest)

    assert xulosa["rasm"]["nusxalandi"] == 1
    assert xulosa["excel"]["yasaldi"] == 1
    assert xulosa["nakladnoy"]["nusxalandi"] == 1
    assert (dest / "KIP-Tarozi Rasm").is_dir()
    assert (dest / "KIP-Tarozi Excel").is_dir()
    assert (dest / "KIP-Tarozi Nakladnoy" / "N-9.pdf").is_file()
