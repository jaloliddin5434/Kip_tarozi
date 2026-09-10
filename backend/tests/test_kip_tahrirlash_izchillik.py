"""0-QISM audit (3-band): kip tahrirlanganda (mahsulot/partiya o'zgarganda)
Dashboard, Statistika va Partiyalar (JSON) HAMDA Excel eksportlari REAL VAQTDA
bazadan o'qiganligi uchun HECH QANDAY qo'shimcha sinxronlashtirishsiz,
darhol to'g'ri bo'lishi kerak — bu yerda aynan shu narsa isbotlanadi (bitta
test ichida: TAHRIRLASHDAN OLDIN va KEYIN solishtiriladi, kesh yo'qligi
kod darajasida ham tasdiqlangan — loyihada `grep -r cache` bo'sh natija beradi)."""

import uuid
from datetime import date, datetime, timezone
from io import BytesIO

from openpyxl import load_workbook

from app.models.foydalanuvchi import Smena
from app.models.mahsulot import Mahsulot
from app.models.partiya import Partiya, PartiyaHolati


def _partiya_yarat(db, mahsulot_id: int, raqami: int) -> Partiya:
    partiya = Partiya(mahsulot_id=mahsulot_id, partiya_raqami=raqami, holati=PartiyaHolati.ochiq)
    db.add(partiya)
    db.commit()
    db.refresh(partiya)
    return partiya


def test_dashboard_statistika_partiyalar_darhol_tuzatilgan_holatni_korsatadi(
    client, db, operator, operator_headers, admin_headers, mahsulot_tola
):
    lint = Mahsulot(kod="lint_izch", nomi="Lint")
    db.add(lint)
    db.commit()
    db.refresh(lint)

    eski_partiya = _partiya_yarat(db, mahsulot_tola.id, 700)
    yangi_partiya = _partiya_yarat(db, lint.id, 701)

    kip = client.post(
        "/api/v1/kiplar",
        json={
            "mijoz_id": str(uuid.uuid4()),
            "partiya_id": eski_partiya.id,
            "ogirlik": 175.5,
            "mahalliy_vaqt": datetime.now(timezone.utc).isoformat(),
        },
        headers=operator_headers,
    ).json()

    # --- TAHRIRLASHDAN OLDIN: Tola=175.5kg/1ta, Lint=0 ---
    dash_oldin = client.get("/api/v1/dashboard", params={"davr": "kunlik"}, headers=admin_headers).json()
    stat_oldin = client.get("/api/v1/statistika/jamlanma", params={"davr": "kunlik"}, headers=admin_headers).json()
    part_eski_oldin = client.get(f"/api/v1/partiyalar", params={"mahsulot_kodi": "tola"}, headers=admin_headers).json()

    def _mahsulot_qatori(javob, kod):
        return next((m for m in javob["mahsulotlar"] if m["mahsulot_kodi"] == kod), None)

    assert _mahsulot_qatori(dash_oldin, "tola")["jami_kg"] == 175.5
    assert _mahsulot_qatori(dash_oldin, "lint_izch") is None or _mahsulot_qatori(dash_oldin, "lint_izch")["soni"] == 0
    assert _mahsulot_qatori(stat_oldin, "tola")["jami_kg"] == 175.5

    eski_partiya_qatori_oldin = next(p for p in part_eski_oldin["items"] if p["id"] == eski_partiya.id)
    assert eski_partiya_qatori_oldin["kip_soni"] == 1
    assert eski_partiya_qatori_oldin["jami_kg"] == 175.5

    # --- TAHRIRLASH: Tola -> Lint ---
    tahrirlash = client.patch(
        f"/api/v1/kiplar/{kip['id']}",
        json={"mahsulot_kodi": "lint_izch", "partiya_raqami": 701, "sabab": "Noto'g'ri mahsulot tanlangan edi"},
        headers=admin_headers,
    )
    assert tahrirlash.status_code == 200, tahrirlash.text

    # --- TAHRIRLASHDAN KEYIN: Tola=0, Lint=175.5kg/1ta — HECH QANDAY qo'shimcha
    # amal (kesh tozalash va h.k.) qilinmadi, faqat qayta so'rov yuborildi. ---
    dash_keyin = client.get("/api/v1/dashboard", params={"davr": "kunlik"}, headers=admin_headers).json()
    stat_keyin = client.get("/api/v1/statistika/jamlanma", params={"davr": "kunlik"}, headers=admin_headers).json()
    part_yangi_keyin = client.get(f"/api/v1/partiyalar", params={"mahsulot_kodi": "lint_izch"}, headers=admin_headers).json()
    part_eski_keyin = client.get(f"/api/v1/partiyalar", params={"mahsulot_kodi": "tola"}, headers=admin_headers).json()

    tola_qator_keyin = _mahsulot_qatori(dash_keyin, "tola")
    assert tola_qator_keyin is None or tola_qator_keyin["jami_kg"] == 0
    assert _mahsulot_qatori(dash_keyin, "lint_izch")["jami_kg"] == 175.5

    stat_tola_keyin = _mahsulot_qatori(stat_keyin, "tola")
    assert stat_tola_keyin is None or stat_tola_keyin["jami_kg"] == 0
    assert _mahsulot_qatori(stat_keyin, "lint_izch")["jami_kg"] == 175.5

    yangi_partiya_qatori = next(p for p in part_yangi_keyin["items"] if p["id"] == yangi_partiya.id)
    assert yangi_partiya_qatori["kip_soni"] == 1
    assert yangi_partiya_qatori["jami_kg"] == 175.5

    eski_partiya_qatori_keyin = next(p for p in part_eski_keyin["items"] if p["id"] == eski_partiya.id)
    assert eski_partiya_qatori_keyin["kip_soni"] == 0
    assert eski_partiya_qatori_keyin["jami_kg"] == 0.0


def test_smena_excel_darhol_tuzatilgan_holatni_korsatadi(
    client, db, operator, operator_headers, admin_headers, mahsulot_tola
):
    """Excel eksporti ham (hisobotlar.py:smena-excel) — real vaqtda bazadan
    o'qiydi, keshsiz. MUHIM: bu endpoint qat'iy 4 ta mahsulot kodini
    (tola/lint/pux/ulyuk) qatorga chiqaradi — shuning uchun test HAQIQIY
    "lint" kodidan foydalanadi, o'zboshimcha kod emas."""
    lint = Mahsulot(kod="lint", nomi="Lint")
    db.add(lint)
    db.commit()
    db.refresh(lint)

    eski_partiya = _partiya_yarat(db, mahsulot_tola.id, 702)
    yangi_partiya = _partiya_yarat(db, lint.id, 703)
    bugun = date.today()

    kip = client.post(
        "/api/v1/kiplar",
        json={
            "mijoz_id": str(uuid.uuid4()),
            "partiya_id": eski_partiya.id,
            "ogirlik": 80.0,
            "mahalliy_vaqt": datetime.now(timezone.utc).isoformat(),
        },
        headers=operator_headers,
    ).json()

    def _excel_qator(mahsulot_nomi):
        javob = client.get(
            "/api/v1/hisobotlar/smena-excel",
            params={"sana": bugun.isoformat(), "smena": operator.smena.value},
            headers=admin_headers,
        )
        assert javob.status_code == 200
        wb = load_workbook(BytesIO(javob.content))
        ws = wb.active
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] == mahsulot_nomi:
                return row
        return None

    assert _excel_qator("Tola")[1:3] == (1, 80.0)

    client.patch(
        f"/api/v1/kiplar/{kip['id']}",
        json={"mahsulot_kodi": "lint", "partiya_raqami": 703, "sabab": "Noto'g'ri mahsulot tanlangan edi"},
        headers=admin_headers,
    )

    tola_qatori_keyin = _excel_qator("Tola")
    assert tola_qatori_keyin[1:3] == (0, 0.0)
    assert _excel_qator("Lint")[1:3] == (1, 80.0)
