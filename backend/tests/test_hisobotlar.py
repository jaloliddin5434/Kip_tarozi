import uuid
from datetime import date, datetime, timezone
from io import BytesIO

from openpyxl import load_workbook

from app.core.security import parolni_hash, token_yarat
from app.models.foydalanuvchi import Foydalanuvchi, Rol, Smena
from app.models.kip import Kip, KipHolati
from app.models.mahsulot import Mahsulot
from app.models.partiya import Partiya, PartiyaHolati


def _partiya_yarat(db, mahsulot_id, raqami):
    partiya = Partiya(mahsulot_id=mahsulot_id, partiya_raqami=raqami, holati=PartiyaHolati.ochiq)
    db.add(partiya)
    db.commit()
    db.refresh(partiya)
    return partiya


def _kip_yarat(db, partiya_id, kip_raqami, ogirlik, smena, sana: date, operator_id):
    vaqt = datetime(sana.year, sana.month, sana.day, 10, 0, tzinfo=timezone.utc)
    kip = Kip(
        mijoz_id=str(uuid.uuid4()),
        partiya_id=partiya_id,
        kip_raqami=kip_raqami,
        ogirlik=ogirlik,
        smena=smena,
        operator_id=operator_id,
        mahalliy_vaqt=vaqt,
        vaqt=vaqt,
        holati=KipHolati.aktiv,
    )
    db.add(kip)
    db.commit()
    return kip


def test_smena_excel_admin_togri_malumot_bilan_generatsiya_qiladi(client, db, admin_headers, operator, mahsulot_tola):
    lint = Mahsulot(kod="lint", nomi="Lint")
    pux = Mahsulot(kod="pux", nomi="Pux")
    ulyuk = Mahsulot(kod="ulyuk", nomi="Ulyuk")
    db.add_all([lint, pux, ulyuk])
    db.commit()

    sana = date.today()
    tola_partiya = _partiya_yarat(db, mahsulot_tola.id, 900)
    lint_partiya = _partiya_yarat(db, lint.id, 901)

    _kip_yarat(db, tola_partiya.id, 1, 100.0, Smena.A, sana, operator.id)
    _kip_yarat(db, tola_partiya.id, 2, 110.0, Smena.A, sana, operator.id)
    _kip_yarat(db, lint_partiya.id, 1, 50.0, Smena.A, sana, operator.id)
    # boshqa smena — hisobotga kirmasligi kerak
    _kip_yarat(db, tola_partiya.id, 3, 999.0, Smena.B, sana, operator.id)

    javob = client.get(
        "/api/v1/hisobotlar/smena-excel", params={"sana": sana.isoformat(), "smena": "A"}, headers=admin_headers
    )
    assert javob.status_code == 200
    assert javob.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert f"Smena_A_{sana.isoformat()}.xlsx" in javob.headers["content-disposition"]

    wb = load_workbook(BytesIO(javob.content))
    ws = wb.active

    qatorlar = {row[0].value: row for row in ws.iter_rows(min_row=2)}

    tola_qatori = qatorlar["Tola"]
    assert tola_qatori[1].value == 2
    assert tola_qatori[2].value == 210.0
    assert tola_qatori[3].value == 105.0

    lint_qatori = qatorlar["Lint"]
    assert lint_qatori[1].value == 1
    assert lint_qatori[2].value == 50.0
    assert lint_qatori[3].value == 50.0

    pux_qatori = qatorlar["Pux"]
    assert pux_qatori[1].value == 0
    assert pux_qatori[2].value == 0.0

    jami_qatori = qatorlar["JAMI"]
    assert jami_qatori[1].value == 3
    assert jami_qatori[2].value == 260.0


def test_smena_excel_tahrirlangan_kip_hisoblanadi_bekor_hisoblanmaydi(
    client, db, admin_headers, operator, mahsulot_tola
):
    sana = date.today()
    partiya = _partiya_yarat(db, mahsulot_tola.id, 940)
    _kip_yarat(db, partiya.id, 1, 100.0, Smena.A, sana, operator.id)
    _kip_yarat(db, partiya.id, 2, 60.0, Smena.A, sana, operator.id).holati = KipHolati.tahrirlangan
    _kip_yarat(db, partiya.id, 3, 999.0, Smena.A, sana, operator.id).holati = KipHolati.bekor_qilingan
    db.commit()

    javob = client.get(
        "/api/v1/hisobotlar/smena-excel", params={"sana": sana.isoformat(), "smena": "A"}, headers=admin_headers
    )
    assert javob.status_code == 200
    ws = load_workbook(BytesIO(javob.content)).active
    qatorlar = {row[0].value: row for row in ws.iter_rows(min_row=2)}
    # aktiv (100) + tahrirlangan (60) = 2 ta / 160 kg; bekor (999) chiqib ketadi
    assert qatorlar["Tola"][1].value == 2
    assert qatorlar["Tola"][2].value == 160.0
    assert qatorlar["JAMI"][1].value == 2
    assert qatorlar["JAMI"][2].value == 160.0


def test_smena_excel_operator_oz_smenasini_olishi_mumkin(client, operator_headers, operator):
    javob = client.get(
        "/api/v1/hisobotlar/smena-excel",
        params={"sana": date.today().isoformat(), "smena": operator.smena.value},
        headers=operator_headers,
    )
    assert javob.status_code == 200


def test_smena_excel_operator_boshqa_smenani_olalmaydi(client, operator_headers):
    javob = client.get(
        "/api/v1/hisobotlar/smena-excel",
        params={"sana": date.today().isoformat(), "smena": "B"},
        headers=operator_headers,
    )
    assert javob.status_code == 403


def test_smena_excel_tokensiz_kira_olmaydi(client):
    javob = client.get("/api/v1/hisobotlar/smena-excel", params={"sana": date.today().isoformat(), "smena": "A"})
    assert javob.status_code == 401


def test_mavsum_jurnali_kunlik_qatorlar_va_jamlanma(client, db, admin_headers, operator, mahsulot_tola):
    boshlanish = date(2025, 9, 1)
    tugash = date(2025, 9, 5)

    partiya = _partiya_yarat(db, mahsulot_tola.id, 500)
    # 1-sentabr: 2 ta kip (100 + 120)
    _kip_yarat(db, partiya.id, 1, 100.0, Smena.A, date(2025, 9, 1), operator.id)
    _kip_yarat(db, partiya.id, 2, 120.0, Smena.B, date(2025, 9, 1), operator.id)
    # 3-sentabr: 1 ta kip (90)
    _kip_yarat(db, partiya.id, 3, 90.0, Smena.A, date(2025, 9, 3), operator.id)
    # 2, 4, 5-sentabr: hech narsa (0 bilan ko'rsatilishi kerak)
    # bekor qilingan kip — hisobga kirmasligi kerak
    _kip_yarat(db, partiya.id, 4, 999.0, Smena.A, date(2025, 9, 4), operator.id).holati = KipHolati.bekor_qilingan
    db.commit()

    javob = client.get(
        "/api/v1/hisobotlar/mavsum-jurnali",
        params={"mahsulot_kodi": "tola", "boshlanish": boshlanish.isoformat(), "tugash": tugash.isoformat()},
        headers=admin_headers,
    )
    assert javob.status_code == 200
    assert javob.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "Mavsum_Jurnali_Tola_2025.xlsx" in javob.headers["content-disposition"]

    ws = load_workbook(BytesIO(javob.content)).active
    qatorlar = {row[0].value: row for row in ws.iter_rows()}

    assert qatorlar["2025-09-01"][2].value == 2
    assert qatorlar["2025-09-01"][3].value == 220.0
    assert qatorlar["2025-09-01"][4].value == 110.0

    assert qatorlar["2025-09-02"][2].value == 0
    assert qatorlar["2025-09-02"][3].value == 0.0

    assert qatorlar["2025-09-03"][2].value == 1
    assert qatorlar["2025-09-03"][3].value == 90.0

    assert qatorlar["2025-09-04"][2].value == 0  # bekor qilingan kip sanalmaydi
    assert qatorlar["2025-09-05"][2].value == 0

    # 5 kun uchun 5 ta kunlik qator bo'lishi kerak
    kun_qatorlari = [r for r in ws.iter_rows() if isinstance(r[0].value, str) and r[0].value.startswith("2025-09-0")]
    assert len(kun_qatorlari) == 5

    oy_jami = qatorlar["Sentabr 2025 — OY JAMI"]
    assert oy_jami[2].value == 3
    assert oy_jami[3].value == 310.0

    mavsum_jami = qatorlar["MAVSUM JAMI"]
    assert mavsum_jami[2].value == 3
    assert mavsum_jami[3].value == 310.0
    assert mavsum_jami[4].value == round(310.0 / 3, 2)


def test_mavsum_jurnali_tahrirlangan_kip_hisoblanadi(client, db, admin_headers, operator, mahsulot_tola):
    partiya = _partiya_yarat(db, mahsulot_tola.id, 502)
    _kip_yarat(db, partiya.id, 1, 100.0, Smena.A, date(2025, 9, 1), operator.id)
    _kip_yarat(db, partiya.id, 2, 80.0, Smena.A, date(2025, 9, 1), operator.id).holati = KipHolati.tahrirlangan
    _kip_yarat(db, partiya.id, 3, 999.0, Smena.A, date(2025, 9, 1), operator.id).holati = KipHolati.bekor_qilingan
    db.commit()

    javob = client.get(
        "/api/v1/hisobotlar/mavsum-jurnali",
        params={"mahsulot_kodi": "tola", "boshlanish": "2025-09-01", "tugash": "2025-09-02"},
        headers=admin_headers,
    )
    assert javob.status_code == 200
    ws = load_workbook(BytesIO(javob.content)).active
    qatorlar = {row[0].value: row for row in ws.iter_rows()}
    # aktiv (100) + tahrirlangan (80) hisoblanadi; bekor (999) yo'q
    assert qatorlar["2025-09-01"][2].value == 2
    assert qatorlar["2025-09-01"][3].value == 180.0
    assert qatorlar["MAVSUM JAMI"][2].value == 2
    assert qatorlar["MAVSUM JAMI"][3].value == 180.0


def test_mavsum_jurnali_oylar_orasida_uzluksiz(client, db, admin_headers, operator, mahsulot_tola):
    partiya = _partiya_yarat(db, mahsulot_tola.id, 501)
    _kip_yarat(db, partiya.id, 1, 50.0, Smena.A, date(2025, 9, 30), operator.id)
    _kip_yarat(db, partiya.id, 2, 70.0, Smena.A, date(2025, 10, 2), operator.id)

    javob = client.get(
        "/api/v1/hisobotlar/mavsum-jurnali",
        params={"mahsulot_kodi": "tola", "boshlanish": "2025-09-29", "tugash": "2025-10-03"},
        headers=admin_headers,
    )
    assert javob.status_code == 200
    ws = load_workbook(BytesIO(javob.content)).active
    qatorlar = {row[0].value: row for row in ws.iter_rows()}

    assert qatorlar["Sentabr 2025 — OY JAMI"][2].value == 1
    assert qatorlar["Sentabr 2025 — OY JAMI"][3].value == 50.0
    assert qatorlar["Oktabr 2025 — OY JAMI"][2].value == 1
    assert qatorlar["Oktabr 2025 — OY JAMI"][3].value == 70.0
    assert qatorlar["MAVSUM JAMI"][2].value == 2
    assert qatorlar["MAVSUM JAMI"][3].value == 120.0


def test_mavsum_jurnali_notogri_mahsulot_404(client, admin_headers):
    javob = client.get(
        "/api/v1/hisobotlar/mavsum-jurnali", params={"mahsulot_kodi": "yoq"}, headers=admin_headers
    )
    assert javob.status_code == 404


def test_mavsum_jurnali_operator_kira_olmaydi(client, operator_headers):
    javob = client.get(
        "/api/v1/hisobotlar/mavsum-jurnali", params={"mahsulot_kodi": "tola"}, headers=operator_headers
    )
    assert javob.status_code == 403


def test_mavsum_jurnali_tokensiz_401(client):
    javob = client.get("/api/v1/hisobotlar/mavsum-jurnali", params={"mahsulot_kodi": "tola"})
    assert javob.status_code == 401


def test_smena_excel_boshqa_rol_kira_olmaydi(client, db):
    foydalanuvchi = Foydalanuvchi(
        ism="Tayyor Test",
        login="tayyor_hisobot_test",
        parol_hash=parolni_hash("parolT"),
        rol=Rol.tayyor_mahsulotlar,
    )
    db.add(foydalanuvchi)
    db.commit()
    db.refresh(foydalanuvchi)

    token = token_yarat(
        {"sub": str(foydalanuvchi.id), "rol": "tayyor_mahsulotlar", "smena": None, "tv": foydalanuvchi.token_versiyasi}
    )
    headers = {"Authorization": f"Bearer {token}"}

    javob = client.get(
        "/api/v1/hisobotlar/smena-excel", params={"sana": date.today().isoformat(), "smena": "A"}, headers=headers
    )
    assert javob.status_code == 403


# ============ smena + mahsulot bo'yicha tor eksport ============


def test_smena_mahsulot_excel_faqat_mos_kiplarni_beradi(client, db, admin_headers, operator, mahsulot_tola):
    lint = Mahsulot(kod="lint", nomi="Lint")
    db.add(lint)
    db.commit()
    db.refresh(lint)

    sana = date.today()
    tola_p = _partiya_yarat(db, mahsulot_tola.id, 930)
    lint_p = _partiya_yarat(db, lint.id, 931)

    _kip_yarat(db, tola_p.id, 1, 100.0, Smena.A, sana, operator.id)
    _kip_yarat(db, tola_p.id, 2, 105.5, Smena.A, sana, operator.id)
    _kip_yarat(db, tola_p.id, 3, 200.0, Smena.B, sana, operator.id)  # boshqa smena
    _kip_yarat(db, lint_p.id, 1, 77.0, Smena.A, sana, operator.id)   # boshqa mahsulot
    bekor = _kip_yarat(db, tola_p.id, 4, 999.0, Smena.A, sana, operator.id)  # bekor qilingan
    bekor.holati = KipHolati.bekor_qilingan
    db.commit()

    javob = client.get(
        "/api/v1/hisobotlar/smena-mahsulot-excel",
        params={"sana": sana.isoformat(), "smena": "A", "mahsulot_kodi": "tola"},
        headers=admin_headers,
    )
    assert javob.status_code == 200
    assert javob.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert f"Smena_A_Tola_{sana.isoformat()}.xlsx" in javob.headers["content-disposition"]

    ws = load_workbook(BytesIO(javob.content)).active
    # kip qatorlari: sarlavha (A1), bo'sh, ustun sarlavhalari, keyin kiplar, keyin JAMI
    kip_qatorlari = [r for r in ws.iter_rows(values_only=True) if isinstance(r[0], int)]
    assert [r[0] for r in kip_qatorlari] == [1, 2]  # faqat Smena A + Tola + aktiv
    assert kip_qatorlari[0][3] == 100.0
    assert kip_qatorlari[1][3] == 105.5

    jami = next(r for r in ws.iter_rows(values_only=True) if r[0] and str(r[0]).startswith("JAMI"))
    assert "2 kip" in jami[0]
    assert jami[3] == 205.5


def test_smena_mahsulot_excel_tahrirlangan_kip_hisoblanadi(client, db, admin_headers, operator, mahsulot_tola):
    sana = date.today()
    partiya = _partiya_yarat(db, mahsulot_tola.id, 932)
    _kip_yarat(db, partiya.id, 1, 100.0, Smena.A, sana, operator.id)
    _kip_yarat(db, partiya.id, 2, 90.0, Smena.A, sana, operator.id).holati = KipHolati.tahrirlangan
    _kip_yarat(db, partiya.id, 3, 999.0, Smena.A, sana, operator.id).holati = KipHolati.bekor_qilingan
    db.commit()

    javob = client.get(
        "/api/v1/hisobotlar/smena-mahsulot-excel",
        params={"sana": sana.isoformat(), "smena": "A", "mahsulot_kodi": "tola"},
        headers=admin_headers,
    )
    assert javob.status_code == 200
    ws = load_workbook(BytesIO(javob.content)).active
    kip_qatorlari = [r for r in ws.iter_rows(values_only=True) if isinstance(r[0], int)]
    assert [r[0] for r in kip_qatorlari] == [1, 2]  # aktiv + tahrirlangan
    jami = next(r for r in ws.iter_rows(values_only=True) if r[0] and str(r[0]).startswith("JAMI"))
    assert "2 kip" in jami[0]
    assert jami[3] == 190.0


def test_smena_mahsulot_excel_operator_oz_smenasi(client, operator_headers, operator, mahsulot_tola):
    javob = client.get(
        "/api/v1/hisobotlar/smena-mahsulot-excel",
        params={"sana": date.today().isoformat(), "smena": operator.smena.value, "mahsulot_kodi": "tola"},
        headers=operator_headers,
    )
    assert javob.status_code == 200


def test_smena_mahsulot_excel_operator_boshqa_smena_403(client, operator_headers, mahsulot_tola):
    javob = client.get(
        "/api/v1/hisobotlar/smena-mahsulot-excel",
        params={"sana": date.today().isoformat(), "smena": "C", "mahsulot_kodi": "tola"},
        headers=operator_headers,
    )
    assert javob.status_code == 403


def test_smena_mahsulot_excel_notogri_mahsulot_404(client, admin_headers):
    javob = client.get(
        "/api/v1/hisobotlar/smena-mahsulot-excel",
        params={"sana": date.today().isoformat(), "smena": "A", "mahsulot_kodi": "yoq"},
        headers=admin_headers,
    )
    assert javob.status_code == 404


def test_smena_mahsulot_excel_tokensiz_401(client):
    javob = client.get(
        "/api/v1/hisobotlar/smena-mahsulot-excel",
        params={"sana": date.today().isoformat(), "smena": "A", "mahsulot_kodi": "tola"},
    )
    assert javob.status_code == 401
