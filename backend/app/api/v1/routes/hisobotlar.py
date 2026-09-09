from datetime import date, timedelta
from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException, Query, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.api.deps import rollarga_ruxsat
from app.core.database import get_db
from app.models.foydalanuvchi import Foydalanuvchi, Rol, Smena
from app.models.kip import HISOBLANADIGAN_HOLATLAR, Kip
from app.models.mahsulot import Mahsulot
from app.models.partiya import Partiya
from app.models.sozlama import Sozlama
from app.services.davr import mavsum_boshlanishi

router = APIRouter(prefix="/hisobotlar", tags=["hisobotlar"])

_MAHSULOT_TARTIBI = ["tola", "lint", "pux", "ulyuk"]

# Mavsum jurnalidagi "boshlanish" sanasini bu sozlama kaliti belgilaydi
# (yo'q yoki noto'g'ri bo'lsa — mavsum_boshlanishi() qoidasiga qaytiladi).
MAVSUM_BOSHI_SOZLAMA_KALITI = "mavsum_boshlanish_sanasi"

_HAFTA_KUNLARI = ["Dushanba", "Seshanba", "Chorshanba", "Payshanba", "Juma", "Shanba", "Yakshanba"]
_OY_NOMLARI = {
    1: "Yanvar", 2: "Fevral", 3: "Mart", 4: "Aprel", 5: "May", 6: "Iyun",
    7: "Iyul", 8: "Avgust", 9: "Sentabr", 10: "Oktabr", 11: "Noyabr", 12: "Dekabr",
}
_OY_FON = PatternFill("solid", fgColor="E8EEF7")
_JAMI_FON = PatternFill("solid", fgColor="D6E4F0")


@router.get("/smena-excel")
def smena_excel(
    sana: date = Query(...),
    smena: Smena = Query(...),
    db: Session = Depends(get_db),
    foydalanuvchi: Foydalanuvchi = Depends(rollarga_ruxsat(Rol.admin, Rol.operator)),
) -> StreamingResponse:
    if foydalanuvchi.rol == Rol.operator and foydalanuvchi.smena != smena:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Faqat o'z smenangiz hisobotini olishingiz mumkin",
        )

    mahsulotlar = db.execute(select(Mahsulot.kod, Mahsulot.nomi)).all()
    nomlar = {kod: nomi for kod, nomi in mahsulotlar}

    qatorlar = db.execute(
        select(Mahsulot.kod, func.count(Kip.id), func.coalesce(func.sum(Kip.ogirlik), 0))
        .join(Partiya, Partiya.mahsulot_id == Mahsulot.id)
        .join(Kip, Kip.partiya_id == Partiya.id)
        .where(Kip.holati.in_(HISOBLANADIGAN_HOLATLAR), Kip.smena == smena, func.date(Kip.vaqt) == sana)
        .group_by(Mahsulot.kod)
    ).all()
    jamlanma = {kod: (soni, float(kg)) for kod, soni, kg in qatorlar}

    wb = Workbook()
    ws = wb.active
    ws.title = f"Smena {smena.value}"

    ws.append(["Mahsulot", "Soni", "Jami kg", "O'rtacha kg"])
    for hujayra in ws[1]:
        hujayra.font = Font(bold=True)

    jami_soni = 0
    jami_kg = 0.0
    for kod in _MAHSULOT_TARTIBI:
        soni, kg = jamlanma.get(kod, (0, 0.0))
        ortacha = round(kg / soni, 2) if soni else 0.0
        ws.append([nomlar.get(kod, kod), soni, round(kg, 2), ortacha])
        jami_soni += soni
        jami_kg += kg

    jami_ortacha = round(jami_kg / jami_soni, 2) if jami_soni else 0.0
    ws.append(["JAMI", jami_soni, round(jami_kg, 2), jami_ortacha])
    for hujayra in ws[ws.max_row]:
        hujayra.font = Font(bold=True)

    for i, kenglik in enumerate([20, 10, 12, 14], start=1):
        ws.column_dimensions[get_column_letter(i)].width = kenglik

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    fayl_nomi = f"Smena_{smena.value}_{sana.isoformat()}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fayl_nomi}"'},
    )


@router.get("/smena-mahsulot-excel")
def smena_mahsulot_excel(
    sana: date = Query(...),
    smena: Smena = Query(...),
    mahsulot_kodi: str = Query(...),
    db: Session = Depends(get_db),
    foydalanuvchi: Foydalanuvchi = Depends(rollarga_ruxsat(Rol.admin, Rol.operator)),
) -> StreamingResponse:
    """Bitta KUN + bitta SMENA + bitta MAHSULOT uchun tor eksport — shu
    kombinatsiyada tortilgan har bir kip alohida qator (kip raqami, vaqt, kg)
    va oxirida JAMI qatori. /smena-excel'dan farqli — u barcha mahsulotlarni
    bitta jamlanma sifatida beradi, bu esa bitta mahsulotning to'liq ro'yxati."""
    if foydalanuvchi.rol == Rol.operator and foydalanuvchi.smena != smena:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Faqat o'z smenangiz hisobotini olishingiz mumkin",
        )

    mahsulot = db.scalar(select(Mahsulot).where(Mahsulot.kod == mahsulot_kodi))
    if mahsulot is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Mahsulot topilmadi")

    qatorlar = db.execute(
        select(Kip.kip_raqami, Kip.vaqt, Kip.ogirlik, Partiya.partiya_raqami)
        .join(Partiya, Partiya.id == Kip.partiya_id)
        .where(
            Partiya.mahsulot_id == mahsulot.id,
            Kip.holati.in_(HISOBLANADIGAN_HOLATLAR),
            Kip.smena == smena,
            func.date(Kip.vaqt) == sana,
        )
        .order_by(Kip.vaqt, Kip.kip_raqami)
    ).all()

    wb = Workbook()
    ws = wb.active
    ws.title = f"Smena {smena.value} — {mahsulot.nomi}"[:31]

    ws.append([f"Smena {smena.value} — {mahsulot.nomi} — {sana.isoformat()}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])

    ws.append(["Kip №", "Partiya №", "Vaqt", "Og'irlik, kg"])
    for hujayra in ws[ws.max_row]:
        hujayra.font = Font(bold=True)

    jami_kg = 0.0
    for kip_raqami, vaqt, ogirlik, partiya_raqami in qatorlar:
        ws.append([kip_raqami, partiya_raqami, vaqt.strftime("%H:%M:%S"), round(float(ogirlik), 2)])
        jami_kg += float(ogirlik)

    ws.append([f"JAMI ({len(qatorlar)} kip)", "", "", round(jami_kg, 2)])
    for hujayra in ws[ws.max_row]:
        hujayra.font = Font(bold=True)
        hujayra.fill = _JAMI_FON

    for i, kenglik in enumerate([18, 12, 12, 14], start=1):
        ws.column_dimensions[get_column_letter(i)].width = kenglik

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    nom_qismi = mahsulot.nomi.replace(" ", "_") or mahsulot.kod
    fayl_nomi = f"Smena_{smena.value}_{nom_qismi}_{sana.isoformat()}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fayl_nomi}"'},
    )


def _mavsum_boshi_sozlamadan(db: Session) -> date | None:
    sozlama = db.get(Sozlama, MAVSUM_BOSHI_SOZLAMA_KALITI)
    if sozlama is None or not sozlama.qiymat:
        return None
    try:
        return date.fromisoformat(sozlama.qiymat.strip())
    except ValueError:
        return None


@router.get("/mavsum-jurnali")
def mavsum_jurnali(
    mahsulot_kodi: str = Query(...),
    boshlanish: date | None = Query(None),
    tugash: date | None = Query(None),
    db: Session = Depends(get_db),
    _: Foydalanuvchi = Depends(rollarga_ruxsat(Rol.admin)),
) -> StreamingResponse:
    """Bitta mahsulot uchun mavsum boshidan bugungacha KUNLIK qatorlar bilan
    jurnal (Excel). Har kun uchun bitta qator — kip tortilmagan kunlar ham
    0 bilan ko'rsatiladi (uzluksiz kalendar). Har oy oxirida "OY JAMI",
    yakunda esa butun mavsum bo'yicha "MAVSUM JAMI" qatori."""
    mahsulot = db.scalar(select(Mahsulot).where(Mahsulot.kod == mahsulot_kodi))
    if mahsulot is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Mahsulot topilmadi")

    bugun = date.today()
    if boshlanish is None:
        boshlanish = _mavsum_boshi_sozlamadan(db) or mavsum_boshlanishi(bugun)
    if tugash is None:
        tugash = bugun
    if tugash < boshlanish:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="tugash sanasi boshlanish sanasidan oldin bo'lishi mumkin emas",
        )

    kun_ustuni = func.date(Kip.vaqt)
    qatorlar = db.execute(
        select(kun_ustuni, func.count(Kip.id), func.coalesce(func.sum(Kip.ogirlik), 0))
        .join(Partiya, Partiya.id == Kip.partiya_id)
        .where(
            Partiya.mahsulot_id == mahsulot.id,
            Kip.holati.in_(HISOBLANADIGAN_HOLATLAR),
            kun_ustuni >= boshlanish,
            kun_ustuni <= tugash,
        )
        .group_by(kun_ustuni)
    ).all()
    kun_jamlanmasi: dict[date, tuple[int, float]] = {}
    for kun, soni, kg in qatorlar:
        k = kun if isinstance(kun, date) else date.fromisoformat(str(kun))
        kun_jamlanmasi[k] = (soni, float(kg))

    wb = Workbook()
    ws = wb.active
    ws.title = f"{mahsulot.nomi} jurnali"[:31]

    ws.append([f"Mavsum jurnali — {mahsulot.nomi}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Davr: {boshlanish.isoformat()} — {tugash.isoformat()}"])
    ws.append([])

    ws.append(["Sana", "Hafta kuni", "Kip soni", "Jami kg", "O'rtacha kg"])
    for hujayra in ws[ws.max_row]:
        hujayra.font = Font(bold=True)

    def jami_qator(sarlavha: str, soni: int, kg: float, fon: PatternFill) -> None:
        ortacha = round(kg / soni, 2) if soni else 0.0
        ws.append([sarlavha, "", soni, round(kg, 2), ortacha])
        for hujayra in ws[ws.max_row]:
            hujayra.font = Font(bold=True)
            hujayra.fill = fon

    mavsum_soni = 0
    mavsum_kg = 0.0
    oy_soni = 0
    oy_kg = 0.0
    joriy_oy = (boshlanish.year, boshlanish.month)

    kun = boshlanish
    while kun <= tugash:
        if (kun.year, kun.month) != joriy_oy:
            jami_qator(f"{_OY_NOMLARI[joriy_oy[1]]} {joriy_oy[0]} — OY JAMI", oy_soni, oy_kg, _OY_FON)
            joriy_oy = (kun.year, kun.month)
            oy_soni = 0
            oy_kg = 0.0

        soni, kg = kun_jamlanmasi.get(kun, (0, 0.0))
        ortacha = round(kg / soni, 2) if soni else 0.0
        ws.append([kun.isoformat(), _HAFTA_KUNLARI[kun.weekday()], soni, round(kg, 2), ortacha])

        oy_soni += soni
        oy_kg += kg
        mavsum_soni += soni
        mavsum_kg += kg
        kun += timedelta(days=1)

    jami_qator(f"{_OY_NOMLARI[joriy_oy[1]]} {joriy_oy[0]} — OY JAMI", oy_soni, oy_kg, _OY_FON)
    jami_qator("MAVSUM JAMI", mavsum_soni, mavsum_kg, _JAMI_FON)

    for i, kenglik in enumerate([14, 14, 10, 12, 14], start=1):
        ws.column_dimensions[get_column_letter(i)].width = kenglik

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    nom_qismi = mahsulot.nomi.replace(" ", "_") or mahsulot.kod
    fayl_nomi = f"Mavsum_Jurnali_{nom_qismi}_{tugash.year}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fayl_nomi}"'},
    )
